"""
Interactive terminal QA Report Builder -- cross-platform (Snowflake / Postgres / Redshift).

Prompts for a BEFORE table and an AFTER table, each on its own connection --
same engine or different engines, any combination of Snowflake, Postgres,
and Redshift. Auto-matches columns by name, lets you confirm the join key,
then compares the two tables and writes the same 6-sheet Excel report as
the Snowflake-only version of this tool.

WHY THIS CAN'T BE A SINGLE SQL JOIN (unlike the Snowflake-to-Snowflake tool):
No database engine can JOIN a table that lives in a different engine. So this
version can't push the whole comparison into one SQL statement. Instead:

  1. Each side computes a per-row checksum SERVER-SIDE (MD5 of the compared
     columns, normalized and concatenated) -- same algorithm on all three
     engines, so a checksum computed on Postgres is directly comparable to
     one computed on Snowflake.
  2. Only (join_key, checksum) pairs -- small, fixed-size regardless of how
     many columns you're comparing -- are pulled back to this machine.
  3. Comparing two lists of (key, checksum) is done with an in-memory
     DuckDB join (":memory:" mode -- nothing is ever written to disk).
     This is where "which rows differ" gets figured out, without ever
     joining or downloading the full tables, and without persisting
     anything locally.
  4. ONLY for the keys whose checksums disagree (or that are missing on one
     side) do we pull full row detail -- still capped/batched, same
     principle as the Snowflake-only tool's missing-record sampling. That
     row detail is then categorized into NULL-transition buckets with
     another in-memory DuckDB query, the same SQL-shaped logic the
     Snowflake-only tool runs directly against Snowflake.

This keeps both the "don't download full tables" and "nothing touches local
disk" safety properties. It is not as cheap as a same-engine SQL join: two
live connections, a checksum pass on each side, and an in-memory DuckDB
reconciliation step. Budget more time than the Snowflake-only tool for large
tables -- DuckDB makes the reconciliation step itself fast even at tens of
millions of rows, but the two independent full-table checksum scans on the
source engines are still the dominant cost, same as any cross-engine
approach.

KNOWN LIMITATION: checksums are only comparable across engines if each
engine's IDEA of a given value's text representation matches after casting
(e.g. a NUMBER's decimal formatting, a TIMESTAMP's default string format).
This script normalizes what it reasonably can (trims whitespace, treats
NULL as a fixed sentinel distinct from any real value, uses one hash
algorithm everywhere) but genuine formatting differences between engines
for numeric/date types can still show up as false "deviations". Treat a
result with many deviations across most/all columns as a sign to check
formatting before assuming the data itself is wrong.

Run with: python qa_report_builder_cross.py
"""

import os
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import duckdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

PER_CATEGORY_SAMPLE = 5   # how many distinct before/after value "sub-patterns" to show per
                           # deviation category (e.g. top 5 most common transformations)
EXAMPLES_PER_SUBPATTERN = 5  # how many real example records (join keys) to show for EACH
                               # of those sub-patterns, not just an aggregate count
CATEGORY_ORDER = ["NULL/empty -> value", "Value -> different value", "Value -> NULL/empty"]
MISSING_RECORDS_CAP = 500      # never pull more than this many missing-record rows, even if
                                # the true missing count is far higher
CHECKSUM_BATCH_SIZE = 10_000    # how many keys go in one WHERE key IN (...) batch when
                                 # pulling full-row detail for mismatched/missing keys.
                                 # Each batch is a full network round-trip, independent of
                                 # how fast the query itself runs -- at a large mismatch
                                 # count (real risk given cross-engine formatting false
                                 # positives), a small batch size means thousands of
                                 # round-trips, each paying real network latency. 10,000
                                 # keeps the IN-list comfortably within any engine's
                                 # practical statement-size limits while cutting round-trip
                                 # count 5x versus a smaller batch.
LARGE_TABLE_WARN_ROWS = 5_000_000

TEXT_TYPE_MARKERS = ("VARCHAR", "STRING", "TEXT", "CHAR", "CHARACTER VARYING", "CHARACTER")

NULL_SENTINEL = "‡‡__QA_NULL_SENTINEL_9f3c7a1e2b8d__‡‡"  # a rare Unicode
                                       # character (double dagger) plus a random-looking token --
                                       # deliberately NOT a NUL byte or other control character,
                                       # which Postgres (and most engines) reject outright inside
                                       # a string literal. Distinct from any realistic real value,
                                       # used inside the server-side checksum expression so NULL never
                                       # collides with the literal text "NULL" in real data


# ── engine adapters ──────────────────────────────────────────────────────────
# Each adapter knows how to: open a connection, run a query and get back
# (columns, rows), quote an identifier, and build the two SQL fragments that
# differ across dialects: casting a column to text, and a NULL-safe
# concat_ws-based row checksum. Everything above this layer is engine-agnostic.

CHECKSUM_ARGS_PER_CONCAT = 90  # PostgreSQL (and Redshift) cap every function call at 100
                                 # arguments -- a flat CONCAT_WS('|', col1, col2, ...) on a
                                 # wide table (400+ columns) blows past that
                                 # in one call. Nesting CONCAT_WS calls in chunks under this
                                 # cap produces the EXACT same joined string as one giant call
                                 # would: CONCAT_WS('|', already-'|'-joined chunk strings)
                                 # reproduces the flat join, so this changes nothing about the
                                 # checksum algorithm or cross-engine comparability -- only how
                                 # the SQL expression is shaped to stay under the limit. 90
                                 # leaves headroom under 100 for the outer wrapping call.


def build_concat_ws_checksum(parts):
    """Builds MD5(CONCAT_WS('|', ...)) safely regardless of column count, by
    nesting CONCAT_WS calls in chunks of CHECKSUM_ARGS_PER_CONCAT. Applied
    identically on every engine (not just the ones with a real argument
    limit) so a checksum stays comparable across engines regardless of which
    side happened to need chunking."""
    if len(parts) <= CHECKSUM_ARGS_PER_CONCAT:
        return f"MD5(CONCAT_WS('|', {', '.join(parts)}))"
    chunks = [parts[i:i + CHECKSUM_ARGS_PER_CONCAT] for i in range(0, len(parts), CHECKSUM_ARGS_PER_CONCAT)]
    chunk_exprs = [f"CONCAT_WS('|', {', '.join(chunk)})" for chunk in chunks]
    return f"MD5(CONCAT_WS('|', {', '.join(chunk_exprs)}))"


class SnowflakeAdapter:
    name = "snowflake"

    def connect(self, env_values, dbname):
        # dbname is unused here -- Snowflake's connection isn't scoped to a
        # single database the way Postgres/Redshift are; the database is
        # only used later to qualify schema/table queries. Accepted for a
        # uniform connect(env_values, dbname) signature across all adapters.
        import snowflake.connector
        print("  (Snowflake -- SSO login; a browser window will open)")
        account = env_or_ask(env_values, "  Account", "SNOWFLAKE_ACCOUNT")
        user = env_or_ask(env_values, "  User (email)", "SNOWFLAKE_USER")
        role = env_or_ask(env_values, "  Role", "SNOWFLAKE_ROLE")
        warehouse = env_or_ask(env_values, "  Warehouse", "SNOWFLAKE_WAREHOUSE")
        conn = snowflake.connector.connect(
            account=account, user=user, authenticator="externalbrowser",
            role=role, warehouse=warehouse,
        )
        return conn

    def cast_text(self, col):
        return f"TO_VARCHAR({col})"

    def checksum_expr(self, cols):
        parts = [f"COALESCE(TRIM({self.cast_text(c)}), '{NULL_SENTINEL}')" for c in cols]
        return build_concat_ws_checksum(parts)

    def large_result_cursor(self, conn):
        # snowflake-connector-python already streams results in batches
        # internally, so the default cursor is fine even for tens of
        # millions of rows.
        return conn.cursor()

    def quote(self, ident):
        return ident  # unquoted identifiers are fine; we always work in uppercase

    def schema_query(self, db, schema, table):
        return (
            f"SELECT column_name, data_type FROM {db}.information_schema.columns "
            f"WHERE UPPER(table_schema) = UPPER('{schema}') AND UPPER(table_name) = UPPER('{table}')"
        )

    def count_query(self, fqn):
        return f"SELECT COUNT(*) FROM {fqn}"


class PostgresAdapter:
    name = "postgres"

    def connect(self, env_values, dbname):
        import psycopg2
        host = env_or_ask(env_values, "  Host", "POSTGRES_HOST")
        port = env_or_ask(env_values, "  Port", "POSTGRES_PORT", "5432")
        user = env_or_ask(env_values, "  User", "POSTGRES_USER")
        password = env_or_ask_password(env_values, "  Password", "POSTGRES_PASSWORD")
        conn = psycopg2.connect(host=host, port=port, dbname=dbname, user=user, password=password)
        # autocommit deliberately NOT set here (default is off, i.e. one
        # long-lived read-only transaction for the whole run) -- a named
        # (server-side) cursor, used below for the checksum fetch, only
        # survives across multiple internal FETCH calls while its
        # DECLARE-ing transaction stays open. Under autocommit=True each
        # statement is its own auto-committed transaction, which would
        # invalidate the cursor between batches. Every query this script
        # runs is a SELECT, so holding one transaction open for the run's
        # duration is safe -- there's nothing to roll back.
        return conn

    def cast_text(self, col):
        return f"CAST({col} AS TEXT)"

    def checksum_expr(self, cols):
        parts = [f"COALESCE(TRIM({self.cast_text(c)}), '{NULL_SENTINEL}')" for c in cols]
        return build_concat_ws_checksum(parts)

    def large_result_cursor(self, conn):
        # A plain (unnamed) psycopg2 cursor makes libpq buffer the ENTIRE
        # result set in memory before Python ever sees a single row --
        # fine for a handful of rows, but a full-table checksum fetch on a
        # table with tens of millions of rows can exhaust available memory
        # before fetchall() is even called, raising "out of memory for
        # query result" from deep inside libpq itself (a real failure seen
        # running this against a 23M-row, 474-column table). A named
        # (server-side) cursor instead asks the server to hold the result
        # set and streams it back in itersize-sized batches on demand,
        # keeping client-side memory bounded regardless of table size.
        import uuid
        cur = conn.cursor(name=f"qa_checksum_cursor_{uuid.uuid4().hex}")
        cur.itersize = 50_000
        return cur

    def quote(self, ident):
        return f'"{ident}"'

    def schema_query(self, db, schema, table):
        return (
            f"SELECT column_name, data_type FROM information_schema.columns "
            f"WHERE table_schema = '{schema.lower()}' AND table_name = '{table.lower()}'"
        )

    def count_query(self, fqn):
        return f"SELECT COUNT(*) FROM {fqn}"


class RedshiftAdapter(PostgresAdapter):
    # Redshift is Postgres-wire-compatible; the only real differences that
    # matter for this script are the driver and the default port.
    name = "redshift"

    def connect(self, env_values, dbname):
        import redshift_connector
        host = env_or_ask(env_values, "  Host", "REDSHIFT_HOST")
        port = env_or_ask(env_values, "  Port", "REDSHIFT_PORT", "5439")
        user = env_or_ask(env_values, "  User", "REDSHIFT_USER")
        password = env_or_ask_password(env_values, "  Password", "REDSHIFT_PASSWORD")
        conn = redshift_connector.connect(host=host, port=int(port), database=dbname, user=user, password=password)
        conn.autocommit = True
        return conn

    def large_result_cursor(self, conn):
        # redshift_connector doesn't support named/server-side cursors the
        # way psycopg2 does, so this falls back to a regular cursor
        # (inherited PostgresAdapter behavior would try a named cursor,
        # which redshift_connector's cursor() doesn't accept). Redshift
        # clusters are typically provisioned with far more memory headroom
        # than a Postgres ODS host or a client laptop, so this class of
        # failure is less likely here -- if it does show up on Redshift,
        # this is the place to add real batching.
        return conn.cursor()


ENGINE_ADAPTERS = {"snowflake": SnowflakeAdapter, "postgres": PostgresAdapter, "redshift": RedshiftAdapter}


class TableSide:
    """One side (BEFORE or AFTER) of the comparison: an engine adapter, a live
    connection, and the fully-qualified table it points at. Every query in
    this script goes through .run() so the rest of the code never touches a
    driver-specific cursor object directly."""

    def __init__(self, label):
        self.label = label
        print(f"\n--- {label}: which engine? ---")
        engine_choice = ask(f"  Engine (snowflake / postgres / redshift)").lower()
        if engine_choice not in ENGINE_ADAPTERS:
            print(f"  '{engine_choice}' is not one of snowflake / postgres / redshift. Exiting.")
            sys.exit(1)
        self.adapter = ENGINE_ADAPTERS[engine_choice]()
        self.db = ask(f"  {label} Database")
        print(f"  Connecting ({self.adapter.name})...")
        self.conn = self.adapter.connect(load_env_file(), self.db)
        print("  Connected.")
        self.schema = ask(f"  {label} Schema")
        self.table = ask(f"  {label} Table")
        self.fqn = f"{self.db}.{self.schema}.{self.table}"
        print(f"  -> {self.fqn} ({self.adapter.name})\n")

    def run(self, sql):
        cur = self.conn.cursor()
        cur.execute(sql)
        cols = [d[0].upper() for d in cur.description]
        # fetchall() already returns a list of tuples on every driver used here
        # (psycopg2, snowflake-connector-python, redshift_connector) -- nothing
        # in this script mutates individual rows, so converting each one to a
        # list here was pure overhead. At tens-of-millions-of-rows scale (the
        # checksum fetch) that's tens of millions of avoidable allocations.
        rows = cur.fetchall()
        cur.close()
        return cols, rows

    def run_large(self, sql):
        """Like run(), but for a query that may return millions of rows --
        currently just the full-table checksum fetch. Uses a streaming
        cursor where the adapter provides one (see each adapter's
        large_result_cursor()), so the client never has to buffer an
        entire huge result set in memory at once before fetching begins."""
        cur = self.adapter.large_result_cursor(self.conn)
        cur.execute(sql)
        # IMPORTANT: do not call cur.fetchall() here. On a psycopg2 NAMED
        # (server-side) cursor, itersize only bounds fetches made through
        # the cursor's iterator protocol -- .fetchall() ignores it and
        # issues a single "FETCH ALL FROM <cursor>", pulling the entire
        # remaining result set into one libpq allocation in one shot. That
        # defeats the whole point of a server-side cursor and is exactly
        # why "out of memory for query result" kept happening even after
        # switching to a named cursor. Iterating the cursor instead makes
        # psycopg2 issue bounded "FETCH FORWARD <itersize>" batches under
        # the hood, so no single allocation ever has to hold more than one
        # batch's worth of rows.
        rows = [row for row in cur]
        # A named cursor's .description is not reliably populated until
        # AFTER the first fetch -- psycopg2 only learns the result columns
        # once an actual FETCH runs against the server-side cursor, not
        # from DECLARE CURSOR alone. Reading it before fetching (as run()'s
        # plain-cursor path safely can) returns None here and crashes.
        cols = [d[0].upper() for d in cur.description]
        cur.close()
        return cols, rows

    def close(self):
        # Postgres connections in this tool run without autocommit (see
        # PostgresAdapter.connect()), so the whole run sits in one
        # long-lived read-only transaction -- commit() just closes it out
        # cleanly. A no-op on engines that don't need this (Snowflake,
        # Redshift under autocommit), since there's nothing pending.
        try:
            self.conn.commit()
        except Exception:
            pass
        self.conn.close()


# ── input helpers ────────────────────────────────────────────────────────────
def ask(prompt, default=None):
    suffix = f" [{default}]" if default else ""
    val = input(f"{prompt}{suffix}: ").strip()
    return val or default


def ask_password(prompt):
    import getpass
    return getpass.getpass(f"{prompt}: ")


ENV_FILE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env")


def load_env_file(path=ENV_FILE_PATH):
    """Reads a local, gitignored KEY=value .env file (one engine's block per
    prefix -- POSTGRES_*, SNOWFLAKE_*, REDSHIFT_*) into a dict. Returns {} if
    the file doesn't exist -- this feature is fully optional and backward
    compatible; every field simply falls back to an interactive prompt when
    not found here. Same manual parser already proven in this session's
    driver scripts, not a new pip dependency."""
    values = {}
    if not os.path.exists(path):
        return values
    with open(path) as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                values[k.strip()] = v.strip()
    return values


def env_or_ask(env_values, prompt, key, default=None):
    if key in env_values and env_values[key]:
        print(f"  (using {key} from .env)")
        return env_values[key]
    return ask(prompt, default)


def env_or_ask_password(env_values, prompt, key):
    if key in env_values and env_values[key]:
        print(f"  (using {key} from .env)")
        return env_values[key]
    return ask_password(prompt)


def fmt(v):
    if v is None:
        return "NULL"
    s = str(v).strip()
    return "NULL" if s == "" else s


def run_both_sides(before_fn, after_fn):
    """Runs one no-argument callable per side concurrently on a 2-thread
    pool and returns (before_result, after_result). BEFORE and AFTER always
    use separate connections to separate engines, so there's no shared
    cursor/connection state between the two calls -- safe to run at the same
    time. This halves wall-clock on the two biggest full-table passes
    (checksum fetch, mismatched/missing row detail pulls) instead of waiting
    on one side to finish before starting the other."""
    with ThreadPoolExecutor(max_workers=2) as ex:
        before_future = ex.submit(before_fn)
        after_future = ex.submit(after_fn)
        return before_future.result(), after_future.result()


# ── schema discovery (uniform across all three engines via information_schema) ──
def describe_columns(side):
    cols, rows = side.run(side.adapter.schema_query(side.db, side.schema, side.table))
    name_i = cols.index("COLUMN_NAME")
    type_i = cols.index("DATA_TYPE")
    out = {}
    for r in rows:
        name = str(r[name_i]).upper()
        raw_type = str(r[type_i])
        type_class = "text" if any(m in raw_type.upper() for m in TEXT_TYPE_MARKERS) else "other"
        out[name] = (raw_type, type_class)
    return out


def match_columns(before_cols, after_cols):
    common = sorted(set(before_cols) & set(after_cols))
    return [(name, before_cols[name][1]) for name in common]


# ── checksum-based row matching (the cross-engine replacement for a SQL JOIN) ──
def fetch_checksums(side, columns, join_key, breakdown_column=None):
    """Returns ({join_key_value: checksum}, breakdown_counts) for every row
    on this side -- breakdown_counts is None if breakdown_column wasn't
    given. This is the only full-table pass on either side: it only ever
    returns the key, a fixed-length hash, and (optionally) one extra
    column per row, never the row's actual column data, so it stays
    bounded by row count, not column count.

    Folding the breakdown column into this SAME query -- instead of a
    separate GROUP BY pass -- avoids a second full-table scan on the
    source engine. That scan is the single most expensive part of the
    whole comparison (real warehouse compute over every compared column,
    for every row), so this is worth far more than any local-side
    optimization: it's a genuine ~2x reduction in server-side work per
    side whenever a breakdown column is used."""
    col_names = [c for c, _ in columns]
    checksum_sql = side.adapter.checksum_expr(col_names)
    select_cols = [join_key, f"{checksum_sql} AS ROW_CHECKSUM"]
    if breakdown_column:
        select_cols.append(breakdown_column)
    _, rows = side.run_large(f"SELECT {', '.join(select_cols)} FROM {side.fqn}")

    if not breakdown_column:
        return {r[0]: r[1] for r in rows}, None

    checksums = {}
    group_counts = {}
    group_keys = {}  # group -> set of join_key values, for a true COUNT(DISTINCT)
    for r in rows:
        key, checksum, group = r[0], r[1], r[2]
        checksums[key] = checksum
        group_counts[group] = group_counts.get(group, 0) + 1
        group_keys.setdefault(group, set()).add(key)
    breakdown_counts = {g: (group_counts[g], len(group_keys[g])) for g in group_counts}
    return checksums, breakdown_counts


def fetch_rows_by_keys(side, columns, join_key, keys):
    """Pull full (compared-column) row data for a specific, bounded list of
    key values -- used only for rows we already know need closer inspection
    (checksum mismatch, or missing on the other side), never for the whole
    table. Batches the IN-list so this stays reasonable even for thousands
    of mismatched keys.

    Returns {key: row_tuple}, where row_tuple holds values in the same
    order as `columns` (SQL guarantees SELECT results come back in the
    requested column order) -- NOT a dict per row. A dict-per-row would
    mean building a wide dict for every single mismatched row on a table
    with hundreds of columns, which adds up fast if the mismatch rate is
    high; callers index into the tuple positionally instead
    (see categorize_with_duckdb's row_tuple() and main()'s missing-record
    assembly)."""
    col_names = [c for c, _ in columns]
    select_list = ", ".join(col_names)
    key_idx = col_names.index(join_key)
    out = {}
    keys = list(keys)
    for i in range(0, len(keys), CHECKSUM_BATCH_SIZE):
        batch = keys[i:i + CHECKSUM_BATCH_SIZE]
        # Standard SQL escaping (doubling embedded single quotes) -- a string
        # join-key value containing a literal quote (e.g. an apostrophe)
        # would otherwise break this query's syntax, the same class of bug
        # NULL_SENTINEL had before it was fixed to avoid NUL bytes.
        quoted = ", ".join(
            f"'{k.replace(chr(39), chr(39) * 2)}'" if isinstance(k, str) else str(k) for k in batch
        )
        _, rows = side.run(f"SELECT {select_list} FROM {side.fqn} WHERE {join_key} IN ({quoted})")
        for r in rows:
            out[r[key_idx]] = r
    return out


def duckdb_conn():
    """A fresh in-memory DuckDB connection. ':memory:' means nothing is ever
    written to disk -- no database file, no persisted temp data."""
    return duckdb.connect(":memory:")


BULK_INSERT_TARGET_PARAMS_PER_BATCH = 200_000  # cap on total bound parameters
                                                 # (rows * col_count) per INSERT,
                                                 # not on row count alone -- see
                                                 # bulk_insert()'s docstring


def bulk_insert(con, table, col_count, rows):
    """Loads Python row tuples into a DuckDB table via large batched
    multi-row INSERT statements, instead of one INSERT per row
    (con.executemany()). At millions of rows, con.executemany() spends most
    of its time on per-call statement-parse/bind overhead rather than the
    actual insert -- batching cuts the number of statement executions by
    4-5 orders of magnitude (e.g. a few hundred calls instead of tens of
    millions for a table with tens of millions of rows) without adding a
    pandas/pyarrow dependency or writing anything to disk.

    Batch size (row count per INSERT) is derived from col_count so the total
    bound parameters per statement stays capped, rather than being a fixed
    row count regardless of table width. A fixed 50,000-row batch is fine
    for the 2-column checksum table, but for a table with hundreds of
    columns of row-detail that would mean millions of parameters in a
    single INSERT -- slow to build and liable to hit driver/statement
    limits. Capping total parameters instead keeps every batch a similar,
    reasonable size regardless of how wide the table being compared is."""
    batch_size = max(1, BULK_INSERT_TARGET_PARAMS_PER_BATCH // col_count)
    placeholder_group = "(" + ", ".join(["?"] * col_count) + ")"
    for i in range(0, len(rows), batch_size):
        batch = rows[i:i + batch_size]
        values_clause = ", ".join([placeholder_group] * len(batch))
        flat_params = [v for row in batch for v in row]
        con.execute(f"INSERT INTO {table} VALUES {values_clause}", flat_params)


def diff_checksums(before_checksums, after_checksums):
    """Reconciles two {key: checksum} maps with an in-memory DuckDB FULL
    OUTER JOIN instead of Python set/dict operations -- meaningfully faster
    at tens of millions of rows. Still only ever touches keys and hashes,
    never real column values, and never persists anything to disk."""
    # DuckDB needs one consistent column type for the join key; keys are cast
    # to text for the join itself, but the ORIGINAL key objects (which may be
    # int, Decimal, str, etc. depending on the source engine) are preserved
    # via this lookup so downstream queries (WHERE key IN (...)) still use
    # the right native type.
    key_lookup = {}
    for k in before_checksums:
        key_lookup[str(k)] = k
    for k in after_checksums:
        key_lookup.setdefault(str(k), k)

    con = duckdb_conn()
    con.execute("CREATE TABLE before_cs(KEY VARCHAR, CHECKSUM VARCHAR)")
    con.execute("CREATE TABLE after_cs(KEY VARCHAR, CHECKSUM VARCHAR)")
    bulk_insert(con, "before_cs", 2, [(str(k), v) for k, v in before_checksums.items()])
    bulk_insert(con, "after_cs", 2, [(str(k), v) for k, v in after_checksums.items()])
    con.execute("""
        CREATE TABLE joined AS
        SELECT COALESCE(b.KEY, a.KEY) AS KEY, b.CHECKSUM AS B_CS, a.CHECKSUM AS A_CS
        FROM before_cs b
        FULL OUTER JOIN after_cs a ON b.KEY = a.KEY
    """)
    exact_match_count = con.execute(
        "SELECT count(*) FROM joined WHERE B_CS IS NOT NULL AND A_CS IS NOT NULL AND B_CS = A_CS"
    ).fetchone()[0]
    missing_in_after_keys = {key_lookup[r[0]] for r in con.execute("SELECT KEY FROM joined WHERE A_CS IS NULL").fetchall()}
    missing_in_before_keys = {key_lookup[r[0]] for r in con.execute("SELECT KEY FROM joined WHERE B_CS IS NULL").fetchall()}
    mismatched_keys = {key_lookup[r[0]] for r in con.execute(
        "SELECT KEY FROM joined WHERE B_CS IS NOT NULL AND A_CS IS NOT NULL AND B_CS != A_CS").fetchall()}
    con.close()
    return missing_in_after_keys, missing_in_before_keys, mismatched_keys, exact_match_count


def categorize_with_duckdb(columns, exact_match_count, mismatched_keys, before_detail, after_detail):
    """Builds the same per-column NULL-transition breakdown + sample
    value-change pairs the Snowflake-only tool computes directly in SQL --
    here run as an in-memory DuckDB query over the (already pulled) row
    detail for every mismatched key, instead of a Python loop.

    KNOWN SIMPLIFICATION: rows whose overall checksum already matched
    (exact_match_count) are folded straight into each column's "exact
    match" tally below, not split between "both NULL" and "same real
    value" -- distinguishing those two would require pulling the real
    column values for rows that are already known to match, which defeats
    the point of using checksums to avoid exactly that. In practice this
    under-counts "both NULL" and over-counts "exact match" very slightly;
    it does not affect deviation counts (which only come from the
    mismatched-checksum subset, computed exactly)."""
    col_names = [c for c, _ in columns]
    col_types = dict(columns)
    total_both = exact_match_count + len(mismatched_keys)

    if not mismatched_keys:
        null_transition = {c: (0, 0, total_both, 0, 0, total_both) for c in col_names}
        diff_counts = {c: 0 for c in col_names}
        return null_transition, diff_counts, {}

    def row_tuple(detail, k):
        # detail[k] is already a raw row tuple in col_names order (see
        # fetch_rows_by_keys), since categorize_with_duckdb is always called
        # with the same `columns` list that built it -- no per-column
        # dict lookup needed here.
        row = detail.get(k)
        if row is None:
            row = (None,) * len(col_names)
        return (str(k),) + tuple(row)

    con = duckdb_conn()
    col_defs = ", ".join(f'"{c}" VARCHAR' for c in col_names)
    con.execute(f"CREATE TABLE before_rows(KEY VARCHAR, {col_defs})")
    con.execute(f"CREATE TABLE after_rows(KEY VARCHAR, {col_defs})")
    # Same batched bulk_insert() as diff_checksums() -- see its docstring.
    bulk_insert(con, "before_rows", len(col_names) + 1, [row_tuple(before_detail, k) for k in mismatched_keys])
    bulk_insert(con, "after_rows", len(col_names) + 1, [row_tuple(after_detail, k) for k in mismatched_keys])

    def empty_and_norm(alias, c, tc):
        col = f'{alias}."{c}"'
        if tc == "text":
            return f"({col} IS NULL OR TRIM({col}) = '')", f"TRIM({col})"
        return f"({col} IS NULL)", col

    # One single query covering every column's breakdown at once -- not a
    # query per column. With many columns, a per-column loop means paying
    # DuckDB's query-parse/plan overhead once per column, which adds up to
    # a long silent wait with no progress shown; a single query pays that
    # overhead once, no matter how many columns are being compared. This
    # mirrors the Snowflake-only tool's own single-query design.
    select_parts = []
    for c in col_names:
        tc = col_types[c]
        b_empty, b_norm = empty_and_norm("b", c, tc)
        a_empty, a_norm = empty_and_norm("a", c, tc)
        select_parts.append(f"sum(CASE WHEN {b_empty} AND {a_empty} THEN 1 ELSE 0 END)")
        select_parts.append(f"sum(CASE WHEN {b_empty} AND NOT {a_empty} THEN 1 ELSE 0 END)")
        select_parts.append(f"sum(CASE WHEN NOT {b_empty} AND NOT {a_empty} AND {b_norm} = {a_norm} THEN 1 ELSE 0 END)")
        select_parts.append(f"sum(CASE WHEN NOT {b_empty} AND NOT {a_empty} AND {b_norm} != {a_norm} THEN 1 ELSE 0 END)")
        select_parts.append(f"sum(CASE WHEN NOT {b_empty} AND {a_empty} THEN 1 ELSE 0 END)")
    row = con.execute(f"""
        SELECT {', '.join(select_parts)}
        FROM before_rows b JOIN after_rows a ON b.KEY = a.KEY
    """).fetchone()

    null_transition = {}
    diff_counts = {}
    idx = 0
    for c in col_names:
        both_null, null_to_val, exact_among_mismatched, val_to_diff, val_to_null = [x or 0 for x in row[idx:idx + 5]]
        idx += 5
        exact = exact_among_mismatched + exact_match_count
        null_transition[c] = (both_null, null_to_val, exact, val_to_diff, val_to_null, total_both)
        # A column has "deviations" if ANY of the three real-difference
        # categories is nonzero -- not just "value changed to a different
        # value". A column that's exclusively NULL/empty->value or
        # value->NULL/empty (zero val_to_diff) is still a real, worth-seeing
        # difference (data enrichment or data loss respectively) -- using
        # val_to_diff alone here previously caused such columns to be
        # mislabeled "0 deviations" and silently skipped for sample-pair
        # fetching, even though their transition-count breakdown (shown
        # regardless, for every column) reported a nonzero count.
        diff_counts[c] = null_to_val + val_to_diff + val_to_null

    deviating = sorted([c for c in diff_counts if diff_counts[c] > 0], key=lambda c: -diff_counts[c])
    print(f"  {len(deviating)} of {len(col_names)} columns have deviations; fetching up to "
          f"{PER_CATEGORY_SAMPLE} sub-patterns per category, each with up to "
          f"{EXAMPLES_PER_SUBPATTERN} real example records, for each...")
    value_patterns = {}
    for i, c in enumerate(deviating, start=1):
        tc = col_types[c]
        b_empty, b_norm = empty_and_norm("b", c, tc)
        a_empty, a_norm = empty_and_norm("a", c, tc)
        # Two-level sample: first find the top PER_CATEGORY_SAMPLE distinct
        # before/after value pairs ("sub-patterns") per category by
        # frequency, then for EACH of those sub-patterns pull up to
        # EXAMPLES_PER_SUBPATTERN real join-key examples of that exact
        # pair -- not just an aggregate count, so a reviewer can look up an
        # actual record. IS NOT DISTINCT FROM (NULL-safe equality) matters
        # here since before_val/after_val can themselves be NULL (e.g. the
        # "Value -> NULL/empty" category).
        rows = con.execute(f"""
            WITH classified AS (
                SELECT b.KEY AS join_key,
                    CASE
                        WHEN {b_empty} AND NOT {a_empty} THEN 'NULL/empty -> value'
                        WHEN NOT {b_empty} AND {a_empty} THEN 'Value -> NULL/empty'
                        ELSE 'Value -> different value'
                    END AS transition_type,
                    {b_norm} AS before_val, {a_norm} AS after_val
                FROM before_rows b JOIN after_rows a ON b.KEY = a.KEY
                WHERE NOT ({b_empty} AND {a_empty})
                  AND NOT (NOT {b_empty} AND NOT {a_empty} AND {b_norm} = {a_norm})
            ),
            top_patterns AS (
                SELECT transition_type, before_val, after_val, count(*) AS cnt
                FROM classified
                GROUP BY 1, 2, 3
                QUALIFY row_number() OVER (PARTITION BY transition_type ORDER BY count(*) DESC) <= {PER_CATEGORY_SAMPLE}
            )
            SELECT tp.transition_type, tp.before_val, tp.after_val, tp.cnt, c.join_key
            FROM top_patterns tp
            JOIN classified c
              ON c.transition_type = tp.transition_type
             AND c.before_val IS NOT DISTINCT FROM tp.before_val
             AND c.after_val IS NOT DISTINCT FROM tp.after_val
            QUALIFY row_number() OVER (
                PARTITION BY tp.transition_type, tp.before_val, tp.after_val ORDER BY c.join_key
            ) <= {EXAMPLES_PER_SUBPATTERN}
            ORDER BY tp.transition_type, tp.cnt DESC, c.join_key
        """).fetchall()
        by_cat = {cat: [] for cat in CATEGORY_ORDER}
        subpatterns_by_key = {}  # (trans, bval, aval) -> the sub-pattern dict, to append examples onto
        for trans, bval, aval, pcnt, join_key_val in rows:
            sp_key = (trans, bval, aval)
            if sp_key not in subpatterns_by_key:
                sp = {"before": bval, "after": aval, "cnt": pcnt, "examples": []}
                subpatterns_by_key[sp_key] = sp
                by_cat.setdefault(trans, []).append(sp)
            subpatterns_by_key[sp_key]["examples"].append(join_key_val)
        both_null, null_to_val, exact_match, val_to_diff, val_to_null, _ = null_transition[c]
        cat_totals = {"NULL/empty -> value": null_to_val, "Value -> different value": val_to_diff,
                      "Value -> NULL/empty": val_to_null}
        value_patterns[c] = {"by_cat": by_cat, "cat_totals": cat_totals, "total_diff": diff_counts[c]}
        print(f"  [{i}/{len(deviating)}] {c}: {diff_counts[c]:,} deviations")

    con.close()
    return null_transition, diff_counts, value_patterns


# ── main ──────────────────────────────────────────────────────────────────────
def main():
    print("=" * 70)
    print("  QA Report Builder -- cross-platform (Snowflake / Postgres / Redshift)")
    print("=" * 70)

    before = TableSide("BEFORE")
    after = TableSide("AFTER")

    print("Fetching schema for both tables...")
    before_cols = describe_columns(before)
    after_cols = describe_columns(after)
    columns = match_columns(before_cols, after_cols)
    before_only = sorted(set(before_cols) - set(after_cols))
    after_only = sorted(set(after_cols) - set(before_cols))

    print(f"  BEFORE columns: {len(before_cols)} | AFTER columns: {len(after_cols)}")
    print(f"  Matched by name (will be compared): {len(columns)}")
    if before_only:
        print(f"  In BEFORE only (not compared): {', '.join(before_only)}")
    if after_only:
        print(f"  In AFTER only (not compared): {', '.join(after_only)}")
    if not columns:
        print("\nNo matching column names between the two tables -- nothing to compare. Exiting.")
        sys.exit(1)

    join_key = ask("\nJoin key column").upper()
    if join_key not in {c for c, _ in columns}:
        print(f"'{join_key}' is not a matched column on both sides. Exiting.")
        sys.exit(1)

    breakdown_column = ask(
        "\nOptional: column to break row counts down by, e.g. tenant_code (press Enter to skip)"
    )
    if breakdown_column:
        breakdown_column = breakdown_column.upper()
        if breakdown_column not in {c for c, _ in columns}:
            print(f"'{breakdown_column}' is not a matched column on both sides -- skipping the breakdown.")
            breakdown_column = None

    print("\nChecking table sizes...")
    (_, before_cnt), (_, after_cnt) = run_both_sides(
        lambda: before.run(before.adapter.count_query(before.fqn)),
        lambda: after.run(after.adapter.count_query(after.fqn)),
    )
    before_total, after_total = before_cnt[0][0], after_cnt[0][0]
    print(f"  BEFORE: {before_total:,} rows | AFTER: {after_total:,} rows")
    if before_total > LARGE_TABLE_WARN_ROWS or after_total > LARGE_TABLE_WARN_ROWS:
        print("  This is a large table. Unlike the same-engine version of this tool, comparing "
              "across engines needs a checksum pass over every row on both sides plus a Python-side "
              "reconciliation step -- budget noticeably more time than a single-engine SQL join.")

    # The breakdown column (if any) is folded into the same checksum query
    # below rather than queried separately -- see fetch_checksums()'s
    # docstring. This avoids a second full-table scan on each side, which
    # is real warehouse compute cost, not just local overhead. BEFORE and
    # AFTER run concurrently (separate connections, separate engines) --
    # this is the single biggest wall-clock win available: the two full
    # table scans no longer wait on each other.
    bd_note = f", broken down by {breakdown_column}" if breakdown_column else ""
    print(f"\nComputing per-row checksums on BEFORE ({before.adapter.name}) and "
          f"AFTER ({after.adapter.name}) concurrently{bd_note}...")
    (before_checksums, before_breakdown), (after_checksums, after_breakdown) = run_both_sides(
        lambda: fetch_checksums(before, columns, join_key, breakdown_column),
        lambda: fetch_checksums(after, columns, join_key, breakdown_column),
    )
    print(f"  BEFORE: {len(before_checksums):,} rows | AFTER: {len(after_checksums):,} rows")

    breakdown_counts = None
    if breakdown_column:
        breakdown_counts = {"column": breakdown_column, "before": before_breakdown, "after": after_breakdown}
        print(f"  {len(set(before_breakdown) | set(after_breakdown))} distinct {breakdown_column} values")

    print("Reconciling checksums with an in-memory DuckDB join...")
    missing_in_after_keys, missing_in_before_keys, mismatched_keys, exact_match_count = diff_checksums(
        before_checksums, after_checksums
    )
    total_both = exact_match_count + len(mismatched_keys)
    present_in_both = total_both
    missing_in_after = len(missing_in_after_keys)
    missing_in_before = len(missing_in_before_keys)
    print(f"  present_in_both={present_in_both:,}  missing_in_after={missing_in_after:,}  "
          f"missing_in_before={missing_in_before:,}  checksum mismatches={len(mismatched_keys):,}")

    print(f"\nPulling full row detail for {len(mismatched_keys):,} mismatched keys "
          f"(batches of {CHECKSUM_BATCH_SIZE})...")
    before_detail, after_detail = run_both_sides(
        lambda: fetch_rows_by_keys(before, columns, join_key, mismatched_keys) if mismatched_keys else {},
        lambda: fetch_rows_by_keys(after, columns, join_key, mismatched_keys) if mismatched_keys else {},
    )

    print("Building per-column NULL-transition breakdown with an in-memory DuckDB query...")
    null_transition, diff_counts, value_patterns = categorize_with_duckdb(
        columns, exact_match_count, mismatched_keys, before_detail, after_detail
    )

    context_cols = [join_key] + [c for c, _ in columns if c != join_key][:6]
    print(f"\nPulling missing-record detail (capped at {MISSING_RECORDS_CAP} each, "
          f"{len(context_cols)} columns of context)...")
    missing_after_sample_keys = list(missing_in_after_keys)[:MISSING_RECORDS_CAP]
    missing_before_sample_keys = list(missing_in_before_keys)[:MISSING_RECORDS_CAP]
    missing_after_detail, missing_before_detail = run_both_sides(
        lambda: fetch_rows_by_keys(before, [(c, dict(columns)[c]) for c in context_cols],
                                    join_key, missing_after_sample_keys) if missing_after_sample_keys else {},
        lambda: fetch_rows_by_keys(after, [(c, dict(columns)[c]) for c in context_cols],
                                    join_key, missing_before_sample_keys) if missing_before_sample_keys else {},
    )
    # missing_after_detail[k] is already a row tuple in context_cols order
    # (fetch_rows_by_keys was called with columns scoped to context_cols) --
    # no per-column lookup needed. .get() with a None-filled fallback (not
    # direct [k] indexing) so a row that vanished between the checksum pass
    # and this later query (a real, if rare, race condition on a live table)
    # can't crash the whole run with a KeyError after all the expensive work
    # is already done -- it just shows as blank context in the report.
    context_width = len(context_cols)
    missing_in_after_rows = [missing_after_detail.get(k, (None,) * context_width) for k in missing_after_sample_keys]
    missing_in_before_rows = [missing_before_detail.get(k, (None,) * context_width) for k in missing_before_sample_keys]

    before.close()
    after.close()
    print("\nDone querying -- building the workbook...")

    build_workbook(
        f"{before.fqn}  ({before.adapter.name})", f"{after.fqn}  ({after.adapter.name})",
        join_key, columns, before_only, after_only,
        before_cols, after_cols, breakdown_counts,
        before_total, after_total, present_in_both, missing_in_after, missing_in_before,
        diff_counts, total_both, null_transition, value_patterns,
        missing_in_after_rows, missing_in_before_rows, context_cols,
        after_table_name=after.table,
    )


# ── excel output (identical style/format to the Snowflake-only tool) ─────────
C = {"navy": "000000", "section": "000000", "bef_hdr": "000000", "bef": "FFFFFF",
     "aft_hdr": "000000", "aft": "FFFFFF", "warn": "FFFFFF", "zebra": "FFFFFF",
     "white": "FFFFFF", "steel": "000000", "sep": "FFFFFF"}


def fn(bold=False, sz=10, col="000000"): return Font(bold=bold, size=sz, color="000000", name="Calibri")
def al(h="left", v="center", wrap=False): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


S = Side(style="thin")
B_THIN = Border(left=S, right=S, top=S, bottom=S)


def cell(ws, r, c, val="", bg=None, fnt=None, aln=None, span=1, nfmt=None):
    o = ws.cell(row=r, column=c, value=val)
    o.font = fnt or fn()
    o.alignment = aln or al()
    o.border = B_THIN
    if nfmt: o.number_format = nfmt
    if span > 1:
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
    return o


def section_hdr(ws, row, label, span, bg=None):
    cell(ws, row, 1, f"  {label}", bg=bg or C["section"], fnt=fn(True, 11, C["white"]), aln=al("left"), span=span)
    ws.row_dimensions[row].height = 22


def col_hdr_row(ws, row, headers, bg):
    for ci, h in enumerate(headers, start=1):
        cell(ws, row, ci, h, bg=bg, fnt=fn(True, 9, C["white"]), aln=al("center"))
    ws.row_dimensions[row].height = 16


def build_workbook(before_fqn, after_fqn, join_key, columns, before_only, after_only,
                    before_cols, after_cols, breakdown_counts,
                    before_total, after_total, present_in_both, missing_in_after, missing_in_before,
                    diff_counts, total_both, null_transition, value_patterns,
                    missing_in_after_rows, missing_in_before_rows, context_cols, after_table_name):
    wb = Workbook()
    wb.remove(wb.active)

    # Sheet 1 -- Schema Comparison
    ws1 = wb.create_sheet("Schema Comparison")
    ws1.sheet_view.showGridLines = False
    cell(ws1, 1, 1, f"Schema Comparison  |  BEFORE ({before_fqn}) vs AFTER ({after_fqn})", bg=C["navy"],
         fnt=fn(True, 12, C["white"]), aln=al("center"), span=6)
    ws1.row_dimensions[1].height = 28
    col_hdr_row(ws1, 3, ["COLUMN", "BEFORE TYPE", "AFTER TYPE", "COMPARED?", "TALLY? (formula)", "NOTE"], C["section"])
    R = 4
    FIRST_DATA_ROW = R
    all_names = sorted(set(before_cols) | set(after_cols))
    for name in all_names:
        b_type = before_cols[name][0] if name in before_cols else "--"
        a_type = after_cols[name][0] if name in after_cols else "--"
        compared = name in {c for c, _ in columns}
        note = "" if compared else ("only in BEFORE" if name not in after_cols else "only in AFTER")
        zebra = C["zebra"] if R % 2 else C["white"]
        cell(ws1, R, 1, name, bg=zebra)
        cell(ws1, R, 2, b_type, bg=zebra, aln=al("center"))
        cell(ws1, R, 3, a_type, bg=zebra, aln=al("center"))
        cell(ws1, R, 4, "YES" if compared else "NO", bg=zebra, aln=al("center"))
        tally_formula = f'=IF(D{R}="NO","NOT COMPARED",IF(B{R}=C{R},"TALLY","TYPE MISMATCH"))'
        o = ws1.cell(row=R, column=5, value=tally_formula)
        o.font = fn(True, 9); o.alignment = al("center"); o.border = B_THIN
        cell(ws1, R, 6, note, bg=zebra)
        R += 1
    LAST_DATA_ROW = R - 1
    R += 1
    summary_formula = (
        f'=CONCATENATE("SUMMARY: ",COUNTIF(E{FIRST_DATA_ROW}:E{LAST_DATA_ROW},"TALLY")'
        f'," of ",COUNTIF(D{FIRST_DATA_ROW}:D{LAST_DATA_ROW},"YES"),'
        f'" compared columns tally exactly on type (formula-verified); any TYPE MISMATCH or"'
        f'" NOT COMPARED rows are worth reviewing before treating the two tables as equivalent."'
        f'," Note: cross-engine type names never match exactly (e.g. Snowflake NUMBER vs Postgres'
        f' numeric) -- TYPE MISMATCH here is expected and not itself a problem when BEFORE/AFTER'
        f' are on different engines; judge by the deviation sheets instead.")'
    )
    o = ws1.cell(row=R, column=1, value=summary_formula)
    o.font = fn(True, 10); o.alignment = al("left"); o.border = B_THIN
    ws1.merge_cells(start_row=R, start_column=1, end_row=R, end_column=6)
    ws1.row_dimensions[R].height = 30
    for ci, w in enumerate([30, 18, 18, 12, 19, 24], start=1):
        ws1.column_dimensions[get_column_letter(ci)].width = w

    # Sheet 2 -- Tenant Counts
    ws2 = wb.create_sheet("Tenant Counts")
    ws2.sheet_view.showGridLines = False
    bd_label = breakdown_counts["column"] if breakdown_counts else "(none provided)"
    cell(ws2, 1, 1, f"Tenant Counts  |  broken down by: {bd_label}", bg=C["navy"],
         fnt=fn(True, 13, C["white"]), aln=al("center"), span=7)
    ws2.row_dimensions[1].height = 26
    col_hdr_row(ws2, 3, ["GROUP", "BEFORE ROWS", "AFTER ROWS", "DIFF (formula)",
                          "BEFORE DISTINCT", "AFTER DISTINCT", "MATCH? (formula)"], C["section"])
    R = 4
    FIRST_DATA_ROW = R
    if breakdown_counts:
        before_map, after_map = breakdown_counts["before"], breakdown_counts["after"]
        for grp in sorted(set(before_map) | set(after_map), key=lambda x: str(x)):
            b_cnt, b_dist = before_map.get(grp, (0, 0))
            a_cnt, a_dist = after_map.get(grp, (0, 0))
            zebra = C["zebra"] if R % 2 else C["white"]
            cell(ws2, R, 1, fmt(grp), bg=zebra)
            cell(ws2, R, 2, b_cnt, bg=C["bef"], nfmt="#,##0", aln=al("center"))
            cell(ws2, R, 3, a_cnt, bg=C["aft"], nfmt="#,##0", aln=al("center"))
            diff_o = ws2.cell(row=R, column=4, value=f"=C{R}-B{R}")
            diff_o.font = fn(); diff_o.alignment = al("center"); diff_o.border = B_THIN
            diff_o.number_format = "+#,##0;-#,##0;0"
            cell(ws2, R, 5, b_dist, bg=zebra, nfmt="#,##0", aln=al("center"))
            cell(ws2, R, 6, a_dist, bg=zebra, nfmt="#,##0", aln=al("center"))
            match_o = ws2.cell(row=R, column=7, value=f'=IF(B{R}=C{R},"MATCH","DIFF")')
            match_o.font = fn(True); match_o.alignment = al("center"); match_o.border = B_THIN
            R += 1
        LAST_DATA_ROW = R - 1
        cell(ws2, R, 1, "TOTAL (all groups)", bg=C["steel"], fnt=fn(True, 10, C["white"]), aln=al("center"))
        tot_b = ws2.cell(row=R, column=2, value=f"=SUM(B{FIRST_DATA_ROW}:B{LAST_DATA_ROW})")
        tot_b.font = fn(True, 10); tot_b.alignment = al("center"); tot_b.border = B_THIN; tot_b.number_format = "#,##0"
        tot_a = ws2.cell(row=R, column=3, value=f"=SUM(C{FIRST_DATA_ROW}:C{LAST_DATA_ROW})")
        tot_a.font = fn(True, 10); tot_a.alignment = al("center"); tot_a.border = B_THIN; tot_a.number_format = "#,##0"
        tot_d = ws2.cell(row=R, column=4, value=f"=C{R}-B{R}")
        tot_d.font = fn(True, 10); tot_d.alignment = al("center"); tot_d.border = B_THIN; tot_d.number_format = "+#,##0;-#,##0;0"
        cell(ws2, R, 5, "", bg=C["steel"]); cell(ws2, R, 6, "", bg=C["steel"])
        tot_chk = ws2.cell(row=R, column=7,
                            value=f'=IF(COUNTIF(G{FIRST_DATA_ROW}:G{LAST_DATA_ROW},"DIFF")=0,"MATCH","DIFF")')
        tot_chk.font = fn(True, 10); tot_chk.alignment = al("center"); tot_chk.border = B_THIN
        R += 1
    else:
        cell(ws2, R, 1, "ALL ROWS", bg=C["white"])
        cell(ws2, R, 2, before_total, bg=C["white"], nfmt="#,##0", aln=al("center"))
        cell(ws2, R, 3, after_total, bg=C["white"], nfmt="#,##0", aln=al("center"))
        cell(ws2, R, 4, after_total - before_total, bg=C["white"], nfmt="+#,##0;-#,##0;0", aln=al("center"))
        cell(ws2, R, 5, "", bg=C["white"]); cell(ws2, R, 6, "", bg=C["white"])
        cell(ws2, R, 7, "MATCH" if before_total == after_total else "DIFF", bg=C["white"], aln=al("center"))
        R += 1
        cell(ws2, R + 1, 1, "No breakdown column was provided for this run -- re-run and answer the "
                             "'column to break row counts down by' prompt (e.g. tenant_code) to see a "
                             "per-group table here instead of one overall row.",
             bg=C["warn"], aln=al("left", wrap=True), span=7)
        ws2.row_dimensions[R + 1].height = 32
    ws2.column_dimensions["A"].width = 26
    for ci in range(2, 8):
        ws2.column_dimensions[get_column_letter(ci)].width = 17

    # Sheet 3 -- Overall Counts
    ws3 = wb.create_sheet("Overall Counts")
    ws3.sheet_view.showGridLines = False
    cell(ws3, 1, 1, "Overall Counts", bg=C["navy"], fnt=fn(True, 13, C["white"]), aln=al("center"), span=2)
    ws3.row_dimensions[1].height = 26
    cell(ws3, 2, 1, "BEFORE", bg=C["bef_hdr"], fnt=fn(True, 9, C["white"]))
    cell(ws3, 2, 2, before_fqn, bg=C["bef"])
    cell(ws3, 3, 1, "AFTER", bg=C["aft_hdr"], fnt=fn(True, 9, C["white"]))
    cell(ws3, 3, 2, after_fqn, bg=C["aft"])
    cell(ws3, 4, 1, "Join key", bg=C["zebra"])
    cell(ws3, 4, 2, join_key, bg=C["zebra"])
    rows = [
        ("BEFORE total rows", before_total), ("AFTER total rows", after_total),
        ("Present in both", present_in_both), ("Missing in AFTER (BEFORE only)", missing_in_after),
        ("Missing in BEFORE (AFTER only)", missing_in_before),
        ("Columns compared", len(columns)),
        ("Columns with 0 deviations", sum(1 for v in diff_counts.values() if v == 0)),
        ("Columns with >0 deviations", sum(1 for v in diff_counts.values() if v > 0)),
        ("Columns only in BEFORE (not compared)", len(before_only)),
        ("Columns only in AFTER (not compared)", len(after_only)),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (label, val) in enumerate(rows, start=6):
        zebra = C["zebra"] if i % 2 else C["white"]
        cell(ws3, i, 1, label, bg=zebra)
        is_num = isinstance(val, int)
        cell(ws3, i, 2, val, bg=zebra, nfmt="#,##0" if is_num else None, aln=al("center") if is_num else al("left"))
    ws3.column_dimensions["A"].width = 40
    ws3.column_dimensions["B"].width = 60

    # Sheet 4 -- Column Mismatch Summary
    ws4 = wb.create_sheet("Column Mismatch Summary")
    ws4.sheet_view.showGridLines = False
    cell(ws4, 1, 1, f"Column Mismatch Summary -- {total_both:,} rows present in both", bg=C["navy"],
         fnt=fn(True, 13, C["white"]), aln=al("center"), span=4)
    ws4.row_dimensions[1].height = 26
    col_hdr_row(ws4, 3, ["COLUMN", "DEVIATIONS", "% OF ROWS", "JUSTIFICATION (fill in after review)"], C["section"])
    R = 4
    for c, cnt in sorted(diff_counts.items(), key=lambda x: -x[1]):
        pct = cnt / total_both if total_both else 0
        zebra = C["zebra"] if R % 2 else C["white"]
        cell(ws4, R, 1, c, bg=zebra)
        cell(ws4, R, 2, cnt, bg=zebra, nfmt="#,##0", aln=al("center"))
        cell(ws4, R, 3, pct, bg=zebra, nfmt="0.00%", aln=al("center"))
        cell(ws4, R, 4, "", bg=zebra)
        R += 1
    ws4.column_dimensions["A"].width = 36
    ws4.column_dimensions["B"].width = 14
    ws4.column_dimensions["C"].width = 12
    ws4.column_dimensions["D"].width = 60

    # Sheet 5 -- Missing Records
    ws5 = wb.create_sheet("Missing Records")
    ws5.sheet_view.showGridLines = False
    span5 = len(context_cols)
    cell(ws5, 1, 1, f"Missing Records -- {missing_in_after:,} missing in AFTER, {missing_in_before:,} missing in BEFORE "
                     f"(sample capped at {MISSING_RECORDS_CAP} each, {span5} columns of context)",
         bg=C["navy"], fnt=fn(True, 12, C["white"]), aln=al("center"), span=span5)
    ws5.row_dimensions[1].height = 30
    section_hdr(ws5, 3, f"Missing in AFTER ({len(missing_in_after_rows)} of {missing_in_after:,} shown)", span5, bg=C["bef_hdr"])
    col_hdr_row(ws5, 4, context_cols, C["bef_hdr"])
    R = 5
    for row in missing_in_after_rows:
        zebra = C["zebra"] if R % 2 else C["white"]
        for ci, val in enumerate(row, start=1):
            cell(ws5, R, ci, fmt(val), bg=zebra)
        R += 1
    R += 1
    section_hdr(ws5, R, f"Missing in BEFORE ({len(missing_in_before_rows)} of {missing_in_before:,} shown)", span5, bg=C["aft_hdr"]); R += 1
    col_hdr_row(ws5, R, context_cols, C["aft_hdr"]); R += 1
    for row in missing_in_before_rows:
        zebra = C["zebra"] if R % 2 else C["white"]
        for ci, val in enumerate(row, start=1):
            cell(ws5, R, ci, fmt(val), bg=zebra)
        R += 1
    for ci in range(1, span5 + 1):
        ws5.column_dimensions[get_column_letter(ci)].width = 24

    # Sheet 6 -- NULL & Value-Change Patterns
    ws6 = wb.create_sheet("NULL & Value-Change Patterns")
    ws6.sheet_view.showGridLines = False
    cell(ws6, 1, 1, "NULL & Value-Change Patterns -- present-in-both rows only", bg=C["navy"],
         fnt=fn(True, 13, C["white"]), aln=al("center"), span=4)
    ws6.row_dimensions[1].height = 26
    R = 2
    for c, _ in columns:
        both_null, null_to_val, exact_match, val_to_diff, val_to_null, total = null_transition[c]
        R += 1
        section_hdr(ws6, R, f"{c} -- {diff_counts[c]:,} deviations of {total:,} rows", 4); R += 1
        col_hdr_row(ws6, R, ["TRANSITION", "COUNT", "% OF ROWS", ""], C["section"]); R += 1
        for label, cnt in [("Both NULL/empty", both_null), ("NULL/empty -> value", null_to_val),
                            ("Exact match", exact_match), ("Value -> different value", val_to_diff),
                            ("Value -> NULL/empty", val_to_null)]:
            pct = (cnt / total) if total else 0
            zebra = C["zebra"] if R % 2 else C["white"]
            cell(ws6, R, 1, label, bg=zebra)
            cell(ws6, R, 2, cnt, bg=zebra, nfmt="#,##0", aln=al("center"))
            cell(ws6, R, 3, pct, bg=zebra, nfmt="0.000%", aln=al("center"))
            cell(ws6, R, 4, "", bg=zebra)
            R += 1
        if c in value_patterns:
            vp = value_patterns[c]
            for cat in CATEGORY_ORDER:
                cat_total = vp["cat_totals"].get(cat, 0)
                cat_subpatterns = vp["by_cat"].get(cat, [])
                R += 1
                if cat_total == 0:
                    section_hdr(ws6, R, f"{cat} -- 0 deviations", 4); R += 1
                    continue
                section_hdr(ws6, R, f"{cat} -- top {len(cat_subpatterns)} sub-pattern(s) "
                                     f"(of up to {PER_CATEGORY_SAMPLE}) out of {cat_total:,} total deviations", 4)
                R += 1
                for sp in cat_subpatterns:
                    cell(ws6, R, 1, f"{fmt(sp['before'])}  ->  {fmt(sp['after'])}", bg=C["white"],
                         fnt=fn(True, 9), aln=al("left", wrap=True), span=3)
                    cell(ws6, R, 4, f"{sp['cnt']:,} rows total", bg=C["white"], fnt=fn(True, 9), aln=al("right"))
                    R += 1
                    col_hdr_row(ws6, R, ["JOIN_KEY", "BEFORE VALUE", "AFTER VALUE",
                                          f"{len(sp['examples'])} example(s) shown"], C["section"])
                    R += 1
                    for i, join_key_val in enumerate(sp["examples"]):
                        zebra = C["zebra"] if i % 2 else C["white"]
                        cell(ws6, R, 1, fmt(join_key_val), bg=zebra)
                        cell(ws6, R, 2, fmt(sp["before"]), bg=zebra, aln=al("left", wrap=True))
                        cell(ws6, R, 3, fmt(sp["after"]), bg=zebra, aln=al("left", wrap=True))
                        cell(ws6, R, 4, "", bg=zebra)
                        R += 1
                    R += 1
        if diff_counts.get(c, 0) > 0:
            cell(ws6, R, 1, "JUSTIFICATION (fill in after review): ",
                 bg=C["warn"], fnt=fn(True, 8, col="000000"), aln=al("left", wrap=True), span=4)
            ws6.row_dimensions[R].height = 45
            R += 1
        R += 1
    for ci, w in enumerate([28, 34, 34, 20], start=1):
        ws6.column_dimensions[get_column_letter(ci)].width = w

    out_path = f"{after_table_name}_cross_qa_report.xlsx"
    wb.save(out_path)
    print(f"\nDone! Saved -> {out_path}")


if __name__ == "__main__":
    main()
